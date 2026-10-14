"""Freeze and normalize the usable STOCKS-26 historical price snapshot."""

import csv
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable


PRICE_COLUMNS = (
    ("AAPL", "USD", 0, 1),
    ("BTC-EUR", "EUR", 2, 3),
    ("EQQQ", "GBX", 4, 5),
    ("MSFT", "USD", 6, 7),
    ("NVDA", "USD", 8, 9),
    ("SPCX", "USD", 10, 11),
    # Active initial benchmark from SETUP: Nasdaq-100 ETF proxy, price return.
    ("QQQ", "USD", 12, 13),
)
GBP_FX_COLUMNS = (80, 81)
USD_FX_COLUMNS = (82, 83)


@dataclass(frozen=True)
class SnapshotObservation:
    ticker: str
    date: date
    price: Decimal
    currency: str
    fx_to_eur: Decimal


@dataclass(frozen=True)
class SnapshotNormalization:
    observations: list[SnapshotObservation]
    errors: list[str]


def normalize_snapshot_rows(rows: Iterable[list[object | None] | tuple[object | None, ...]]) -> SnapshotNormalization:
    materialized_rows = list(rows)
    fx_by_currency = {
        "GBP": _fx_values(materialized_rows, GBP_FX_COLUMNS),
        "USD": _fx_values(materialized_rows, USD_FX_COLUMNS),
    }
    observations: list[SnapshotObservation] = []
    errors: list[str] = []
    for ticker, currency, date_column, price_column in PRICE_COLUMNS:
        for row in materialized_rows:
            raw_date, raw_price = _get(row, date_column), _get(row, price_column)
            if raw_date in (None, "") and raw_price in (None, ""):
                continue
            try:
                observed_on = _as_date(raw_date)
                price = _as_decimal(raw_price)
                if price <= 0:
                    raise ValueError("prezzo non positivo")
                fx_currency = "GBP" if currency == "GBX" else currency
                fx_to_eur = Decimal("1") if currency == "EUR" else _latest_fx(observed_on, fx_by_currency[fx_currency], fx_currency)
                observations.append(SnapshotObservation(ticker, observed_on, price, currency, fx_to_eur))
            except (InvalidOperation, ValueError) as error:
                errors.append(f"{ticker} {_date_label(raw_date)}: {error}")
    observations.sort(key=lambda item: (item.ticker, item.date))
    return SnapshotNormalization(observations, errors)


def extract_xlsx(source: Path) -> SnapshotNormalization:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["_PRICE_HISTORY"]
        return normalize_snapshot_rows(sheet.iter_rows(min_row=6, max_col=84, values_only=True))
    finally:
        workbook.close()


def write_normalized_csv(snapshot: SnapshotNormalization, destination: Path) -> None:
    if snapshot.errors:
        raise ValueError("Snapshot non importabile: " + "; ".join(snapshot.errors[:5]))
    with destination.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["ticker", "date", "price", "currency", "fx_to_eur"])
        writer.writeheader()
        for item in snapshot.observations:
            writer.writerow(
                {
                    "ticker": item.ticker,
                    "date": item.date.isoformat(),
                    "price": item.price,
                    "currency": item.currency,
                    "fx_to_eur": item.fx_to_eur,
                }
            )


def _fx_values(rows: list[list[object | None] | tuple[object | None, ...]], columns: tuple[int, int]) -> dict[date, Decimal]:
    result: dict[date, Decimal] = {}
    for row in rows:
        raw_date, raw_rate = _get(row, columns[0]), _get(row, columns[1])
        if raw_date in (None, "") or raw_rate in (None, ""):
            continue
        try:
            observed_on, rate = _as_date(raw_date), _as_decimal(raw_rate)
            if rate <= 0:
                continue
            result[observed_on] = rate
        except (InvalidOperation, ValueError):
            continue
    return result


def _latest_fx(observed_on: date, rates: dict[date, Decimal], currency: str) -> Decimal:
    eligible_dates = [rate_date for rate_date in rates if rate_date <= observed_on]
    if not eligible_dates:
        raise ValueError(f"cambio {currency}/EUR non disponibile")
    return rates[max(eligible_dates)]


def _as_date(value: object | None) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (float, int)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    if isinstance(value, str):
        return date.fromisoformat(value[:10])
    raise ValueError("data non valida")


def _as_decimal(value: object | None) -> Decimal:
    if value in (None, ""):
        raise ValueError("prezzo mancante")
    return Decimal(str(value))


def _get(row: list[object | None] | tuple[object | None, ...], index: int) -> object | None:
    return row[index] if len(row) > index else None


def _date_label(value: object | None) -> str:
    try:
        return _as_date(value).isoformat()
    except ValueError:
        return "data non valida"
